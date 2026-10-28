from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt


def write_final_report(
    customers: pd.DataFrame,
    results: pd.DataFrame,
    importance: pd.DataFrame,
    scenario: pd.DataFrame,
    config: dict[str, Any],
    path: str | Path,
) -> None:
    churn_rate = float(customers["churn"].mean())
    best_all = results.loc[results["feature_set"] == "all_features"].iloc[0]
    best_demo = results.loc[results["feature_set"] == "demographics_only"].iloc[0]
    best_behavior = results.loc[
        results["feature_set"] == "behavior_engagement_support"
    ].iloc[0]
    top_features = importance.head(10)
    scenario_row = scenario.iloc[0]
    dataset_label = config.get("dataset_label", "UNSPECIFIED DATA SOURCE")
    report = f"""# FreshBasket Customer Churn Prediction

> **{dataset_label}**

## Executive summary

- Customers with usable profile and churn labels: **{len(customers):,}**.
- Observed churn-label rate in the joined analysis population: **{churn_rate:.1%}**.
- Label window configured: **{config['label_window']['start']} through {config['label_window']['end']}**.
- Best all-feature temporal-validation model: **{best_all['model']}** (PR-AUC {best_all['pr_auc']:.3f}, F1 {best_all['f1']:.3f}).
- Demographics-only best PR-AUC: **{best_demo['pr_auc']:.3f}** ({best_demo['model']}).
- Behaviour, engagement, and support only best PR-AUC: **{best_behavior['pr_auc']:.3f}** ({best_behavior['model']}).

## Validation and leakage controls

The latest signup cohort(s) are held out for temporal validation. No random split is used.
Monthly activity from **{config['label_window']['start']} onward is excluded** from features.
The validation metrics below reflect the held-out cohort and are not in-sample fit metrics.

{results.to_markdown(index=False)}

## Leading model cues

Global permutation importance is measured on the temporal-validation cohort using PR-AUC.
It indicates predictive association, not causation.

{top_features.to_markdown(index=False) if not top_features.empty else 'No importance data generated.'}

Customer-level `key_model_cues` are the globally important feature values for each scored
customer; they are descriptive cues, not individualized causal explanations.

## Retention scenario (assumption-based)

The scenario contacts the highest-risk **{scenario_row['contact_fraction_assumption']:.0%}** of
scored validation customers and assumes an absolute retention lift of
**{scenario_row['assumed_absolute_retention_lift']:.0%}** per contacted customer.
Estimated incremental retained customers: **{scenario_row['estimated_incremental_retained_customers']:.1f}**.
Estimated net value under configured cost/value assumptions:
**{scenario_row['estimated_net_value']:.2f}**.

This is a what-if simulation only. It is not a causal estimate or a proven program impact.
A randomized holdout or other credible causal evaluation is required before claiming impact.
Scenario inputs are in `retention_scenario.csv`.

## Recommendations

1. Use the temporal-validation results and precision/recall trade-off to select an operational
   contact threshold; the current 0.5 classification threshold is a reporting default.
2. Pilot outreach with an untreated randomized control group, tracking incremental retention,
   redemption, contact cost, and margin.
3. Use feature-importance cues to formulate hypotheses for the pilot, not to infer causality.
4. Review the generated quality-check CSVs with data owners, especially missing customer-months,
   duplicate keys, negative spend, zero transactions, casing differences, and IQR outliers.
5. Revalidate on a later label window before production use; do not score with post-window data.

## Generated files

- `processed/cleaned_modeling_dataset.csv`: joined, label-window-safe customer features.
- `outputs/validation_results.csv`: model and feature-set metrics.
- `outputs/customer_predictions.csv`: out-of-time validation predictions only.
- `outputs/feature_importance.csv`: validation permutation importance.
- `outputs/retention_scenario.csv`: explicit assumptions and simulated impact.
- `outputs/quality_*.csv`, `outputs/*churn*.png`, `outputs/feature_importance.png`.
"""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(report, encoding="utf-8")


def create_presentation(report_data: dict[str, Any], path: str | Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    def add_slide(title: str, bullets: list[str], notes: str) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.clear()
        for index, text in enumerate(bullets):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = text
            paragraph.level = 0
            paragraph.font.size = Pt(24)
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = notes

    results = report_data["results"]
    best = results.loc[results["feature_set"] == "all_features"].iloc[0]
    scenario = report_data["scenario"].iloc[0]
    importance = report_data["importance"].head(5)
    cues = ", ".join(importance["feature"].astype(str)) if len(importance) else "Not available"
    dataset_label = report_data.get("dataset_label", "UNSPECIFIED DATA SOURCE")
    add_slide(
        "FreshBasket Loyalty Churn — SYNTHETIC DEMO",
        [dataset_label, "April–June 2024 label window", "Decision support, not causal proof"],
        "Suggested timing: 30 seconds. Introduce the business question and distinguish prediction from intervention impact.",
    )
    add_slide(
        "Business question and data",
        [
            dataset_label,
            f"{len(report_data['customers']):,} customers in the joined analysis population",
            f"Observed label rate: {report_data['customers']['churn'].mean():.1%}",
            "Three inputs: profile, monthly activity, and churn label tables",
        ],
        "Suggested timing: 1 minute. Explain the three tables and the need to inspect their actual columns before mapping.",
    )
    add_slide(
        "Quality and leakage controls",
        [
            "Quality checks cover missingness, duplicate IDs, missing months, spend, transactions, casing, and outliers",
            "Monthly features use activity before 1 April 2024 only",
            f"Temporal holdout: latest signup cohort(s) {best['validation_cohorts']}",
        ],
        "Suggested timing: 1 minute. Point to quality CSVs and discuss any findings after the actual run.",
    )
    add_slide(
        "Validation approach",
        [
            f"Selected all-feature model: {best['model']}",
            f"PR-AUC {best['pr_auc']:.3f} | Precision {best['precision']:.3f} | Recall {best['recall']:.3f}",
            f"F1 {best['f1']:.3f} | Accuracy {best['accuracy']:.3f}",
        ],
        "Suggested timing: 2 minutes. Describe the temporal split and explain why PR-AUC is useful for imbalanced labels.",
    )
    ablation = results.pivot(index="feature_set", columns="model", values="pr_auc")
    add_slide(
        "Ablation: feature groups",
        [f"{row}: " + ", ".join(f"{key} PR-AUC {value:.3f}" for key, value in values.items() if pd.notna(value))
         for row, values in ablation.iterrows()],
        "Suggested timing: 1.5 minutes. Compare demographics-only, behavioural/engagement/support, and all-feature results.",
    )
    add_slide(
        "What the model learned",
        [f"Leading validation permutation-importance cues: {cues}",
         "Importance indicates predictive association, not a causal driver"],
        "Suggested timing: 1.5 minutes. Show feature_importance.png and discuss limits of predictive explanation.",
    )
    add_slide(
        "Retention scenario",
        [
            f"Contacted customers: {scenario['customers_contacted']:,}",
            f"Assumed lift: {scenario['assumed_absolute_retention_lift']:.0%}; estimated retained: {scenario['estimated_incremental_retained_customers']:.1f}",
            f"Estimated net value: {scenario['estimated_net_value']:.2f}",
            "Illustrative assumptions only; not a causal impact estimate",
        ],
        "Suggested timing: 1.5 minutes. State the contact and treatment-effect assumptions and why this is not evidence of impact.",
    )
    add_slide(
        "Recommended action",
        [
            "Pilot risk-ranked outreach with randomized treatment and control groups",
            "Track incremental retention, contribution margin, redemption, and contact cost",
            "Refresh validation when a later churn-label window becomes available",
        ],
        "Suggested timing: 1 minute. Close with a measurable controlled pilot and next validation cycle.",
    )
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(target)
